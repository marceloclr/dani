"""Planilha única do projeto (ADR-29): public/modelos/obra-dani.xlsx.

Uma aba por natureza dos dados, com listas suspensas, cabeçalho congelado e o sobrado de exemplo preenchido.
O app lê as abas pelo nome e as colunas pelo cabeçalho (src/planilha/ler.ts); a aba LEIA-ME não é lida.

Uso: ~/.local/venv-ifc/bin/python tools/gerar_planilha_modelo.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path("public/modelos/obra-dani.xlsx")

# Papel e Tinta + marca da Daniella (grafite e dourado)
GRAFITE, DOURADO, PAPEL, CARTA, TINTA_2, LINHA = "2C2C2C", "B88848", "F3F1EC", "FFFEFC", "4A5058", "D9D4CB"
FONTE = "IBM Plex Sans"

MUNICIPIOS = [
    "Fortaleza", "Aquiraz", "Cascavel", "Caucaia", "Chorozinho", "Eusébio", "Guaiúba", "Horizonte", "Itaitinga", "Maracanaú",
    "Maranguape", "Pacajus", "Pacatuba", "Paracuru", "Paraipaba", "Pindoretama", "São Gonçalo do Amarante", "São Luís do Curu",
    "Trairi", "Outro município do Ceará",
]
CATEGORIAS = [
    "terreno", "fundacao", "estrutura", "alvenaria", "laje", "cobertura", "instalacoes", "reboco", "esquadrias", "revestimento",
    "pintura", "loucas", "paisagismo", "limpeza", "entrega",
]
LISTAS = {
    "municipio": MUNICIPIOS,
    "categoria": CATEGORIAS,
    "acao": ["construção", "acabamento", "instalação", "temporário", "remoção"],
    "modo": ["incluir", "excluir"],
    "cena": ["terreno", "sobre a obra", "só a voz"],
    "recorte": ["IA", "fundo verde"],
    "cobertura": ["duas águas", "uma água", "plana"],
    "formato": ["vertical 9:16", "horizontal 16:9", "quadrado 1:1"],
    "qualidade": ["máxima", "normal"],
    "luz": ["Dia", "Nascer", "Entardecer", "Noite", "Ciclo"],
    "animacao": ["Progressivo", "Aparecimento", "Fade-in", "Crescimento", "Por fases"],
    "aparencia": ["Realista", "Técnica"],
    "simnao": ["sim", "não"],
    "fps": ["30", "24"],
}

borda = Border(bottom=Side(style="thin", color=LINHA))
cab_fill = PatternFill("solid", fgColor=GRAFITE)
cab_font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
campo_font = Font(name=FONTE, bold=True, color=GRAFITE, size=10)
nota_font = Font(name=FONTE, italic=True, color=TINTA_2, size=9)
corpo_font = Font(name=FONTE, size=10)

wb = Workbook()
wb.remove(wb.active)
lst = wb.create_sheet("Listas")
lst.sheet_state = "hidden"
intervalos: dict[str, str] = {}
for c, (nome, valores) in enumerate(LISTAS.items(), start=1):
    col = get_column_letter(c)
    lst.cell(1, c, nome)
    for i, v in enumerate(valores, start=2):
        lst.cell(i, c, v)
    intervalos[nome] = f"Listas!${col}$2:${col}${len(valores) + 1}"


def lista(ws, nome: str, alvo: str) -> None:
    dv = DataValidation(type="list", formula1=intervalos[nome], allow_blank=True, showErrorMessage=True,
                        errorTitle="Valor fora da lista", error="Escolha um valor da lista.")
    ws.add_data_validation(dv)
    dv.add(alvo)


def cabecalho(ws, colunas: list[tuple[str, int]]) -> None:
    for c, (titulo, largura) in enumerate(colunas, start=1):
        cel = ws.cell(1, c, titulo)
        cel.fill, cel.font = cab_fill, cab_font
        cel.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = largura
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def ficha(ws, linhas: list[tuple[str, object, str, str | None]]) -> None:
    """Aba de campo e valor: (campo, valor, observação, lista)."""
    cabecalho(ws, [("campo", 34), ("valor", 46), ("observação", 70)])
    for r, (campo, valor, nota, nome_lista) in enumerate(linhas, start=2):
        ws.cell(r, 1, campo).font = campo_font
        v = ws.cell(r, 2, valor)
        v.font = corpo_font
        if isinstance(valor, date):
            v.number_format = "DD/MM/YYYY"
        ws.cell(r, 3, nota).font = nota_font
        for c in range(1, 4):
            ws.cell(r, c).border = borda
        if nome_lista:
            lista(ws, nome_lista, f"B{r}")


def tabela(ws, colunas: list[tuple[str, int]], linhas: list[list[object]], listas: dict[int, str], datas: set[int] = frozenset(), ate: int = 400) -> None:
    cabecalho(ws, colunas)
    for r, linha in enumerate(linhas, start=2):
        for c, valor in enumerate(linha, start=1):
            cel = ws.cell(r, c, valor if valor != "" else None)
            cel.font, cel.border = corpo_font, borda
            if c in datas:
                cel.number_format = "DD/MM/YYYY"
    for c, nome_lista in listas.items():
        col = get_column_letter(c)
        lista(ws, nome_lista, f"{col}2:{col}{ate}")
    for c in datas:
        col = get_column_letter(c)
        for r in range(len(linhas) + 2, ate + 1):
            ws[f"{col}{r}"].number_format = "DD/MM/YYYY"


# ---------- LEIA-ME ----------
leia = wb.create_sheet("LEIA-ME", 0)
leia.sheet_view.showGridLines = False
leia.column_dimensions["A"].width = 4
leia.column_dimensions["B"].width = 110
leia["B2"] = "DANIELLA POMPEU · ENGENHARIA QUE TRANSFORMA"
leia["B2"].font = Font(name=FONTE, bold=True, color=DOURADO, size=13)
leia["B3"] = "Planilha da obra: tudo o que o app precisa para gerar o vídeo e o documento"
leia["B3"].font = Font(name=FONTE, bold=True, color=GRAFITE, size=16)
TEXTO = [
    "",
    "Como usar",
    "1. Preencha as abas abaixo (o exemplo é o sobrado de demonstração: troque pelos dados da sua obra).",
    "2. No app, passo Carregar: envie esta planilha e, em cartões separados, o IFC (opcional), os vídeos das falas e as fotos.",
    "3. Passo Conferir: veja os avisos e a prévia. Passo Gerar: baixe o vídeo MP4 e o documento em PDF e DOCX.",
    "",
    "Abas",
    "Obra: dados da obra e da responsável técnica. Data de referência = dia do status do documento (vazio = hoje).",
    "Modelo: só quando não houver IFC; o app gera uma casa com estas medidas.",
    "Cronograma: uma linha por etapa. Datas em DD/MM/AAAA; avanço em % (ex.: 60%). Colunas *_real e avanço são opcionais.",
    "Vínculos: opcional. Corrige elemento por elemento (GUID do IFC) qual etapa o constrói.",
    "Falas: um vídeo da engenheira por linha, na ordem do vídeo final. O nome do arquivo deve ser igual ao enviado.",
    "Fotos: fotos reais da obra, com data e etapa (ID da aba Cronograma).",
    "Vídeo: formato, duração, qualidade e luz do vídeo.",
    "Documento: título, destinatário, observações e seções do PDF/DOCX.",
    "",
    "Os campos com seta têm lista de valores. Não renomeie as abas nem os cabeçalhos.",
    "Manual completo: https://marceloclr.github.io/dani/manual.html",
]
for i, t in enumerate(TEXTO, start=5):
    c = leia.cell(i, 2, t)
    c.font = Font(name=FONTE, bold=t in ("Como usar", "Abas"), color=GRAFITE if t in ("Como usar", "Abas") else TINTA_2, size=11)

# ---------- Obra ----------
ficha(wb.create_sheet("Obra"), [
    ("Nome da obra", "Sobrado de exemplo", "Aparece na capa do documento e no nome dos arquivos.", None),
    ("Proprietário", "Família Exemplo", "", None),
    ("Endereço", "Rua Exemplo, 100 – Fortaleza/CE", "", None),
    ("Município", "Fortaleza", "Define os feriados (dias úteis) e o sol, quando o IFC não traz o local.", "municipio"),
    ("Responsável técnica", "Daniella Pompeu", "", None),
    ("CREA", "", "Ex.: CE-000000", None),
    ("Data de referência", date(2026, 8, 15), "Status mostrado no documento. Vazio = data de hoje.", None),
    ("Arquivo IFC", "sobrado-exemplo.ifc", "Nome do arquivo enviado no cartão Projeto IFC. Vazio = casa gerada pela aba Modelo.", None),
    ("Rumo da fachada frontal (graus)", 70, "Para onde a frente da casa olha (0 = norte, 90 = leste). Vazio = o do IFC.", None),
    ("Descrição", "Sobrado de dois pavimentos com três quartos.", "Texto curto para o documento.", None),
])

# ---------- Modelo ----------
ficha(wb.create_sheet("Modelo"), [
    ("Largura do terreno (m)", 12, "Usado só sem IFC.", None),
    ("Comprimento do terreno (m)", 30, "", None),
    ("Área construída (m²)", 120, "", None),
    ("Pavimentos", 1, "1 ou 2.", None),
    ("Pé-direito (m)", 2.8, "", None),
    ("Cobertura", "duas águas", "", "cobertura"),
])

# ---------- Cronograma ----------
def d(s: str) -> date | str:
    if not s:
        return ""
    dd, mm, aa = s.split("/")
    return date(int(aa), int(mm), int(dd))


ETAPAS = [
    # id, nome, início, fim, categoria, pavimento, início real, fim real, avanço
    ("PRE-01", "Serviços preliminares e locação", "02/03/2026", "13/03/2026", "terreno", "", "02/03/2026", "13/03/2026", "100%"),
    ("FUN-01", "Fundação (sapatas, baldrames e contrapiso)", "14/03/2026", "12/04/2026", "fundacao", "", "16/03/2026", "18/04/2026", "100%"),
    ("EST-T", "Estrutura do térreo e escada", "13/04/2026", "05/05/2026", "estrutura", "Térreo", "20/04/2026", "12/05/2026", "100%"),
    ("ALV-T", "Alvenaria do térreo", "29/04/2026", "24/05/2026", "alvenaria", "Térreo", "06/05/2026", "30/05/2026", "100%"),
    ("LAJ-T", "Laje do pavimento superior", "20/05/2026", "05/06/2026", "laje", "Térreo", "28/05/2026", "14/06/2026", "100%"),
    ("EST-S", "Estrutura do pavimento superior", "06/06/2026", "24/06/2026", "estrutura", "Pavimento superior", "15/06/2026", "03/07/2026", "100%"),
    ("ALV-S", "Alvenaria do pavimento superior", "20/06/2026", "15/07/2026", "alvenaria", "Pavimento superior", "30/06/2026", "28/07/2026", "100%"),
    ("LAJ-S", "Laje de forro", "12/07/2026", "26/07/2026", "laje", "Pavimento superior", "25/07/2026", "", "70%"),
    ("COB-01", "Telhado e oitões", "27/07/2026", "20/08/2026", "cobertura", "", "", "", ""),
    ("INS-01", "Instalações hidráulicas e elétricas", "01/07/2026", "05/09/2026", "instalacoes", "", "08/07/2026", "", "35%"),
    ("REB-01", "Chapisco e reboco", "21/08/2026", "20/09/2026", "reboco", "", "", "", ""),
    ("ESQ-01", "Esquadrias", "21/09/2026", "10/10/2026", "esquadrias", "", "", "", ""),
    ("PIS-01", "Revestimento de pisos", "15/09/2026", "12/10/2026", "revestimento", "", "", "", ""),
    ("PIN-01", "Pintura", "11/10/2026", "05/11/2026", "pintura", "", "", "", ""),
    ("LOU-01", "Louças e metais", "25/10/2026", "08/11/2026", "loucas", "", "", "", ""),
    ("PAI-01", "Paisagismo", "01/11/2026", "18/11/2026", "paisagismo", "", "", "", ""),
    ("LIM-01", "Limpeza final", "19/11/2026", "23/11/2026", "limpeza", "", "", "", ""),
    ("ENT-01", "Vistoria e entrega", "24/11/2026", "26/11/2026", "entrega", "", "", "", ""),
]
tabela(
    wb.create_sheet("Cronograma"),
    [("id", 10), ("nome", 44), ("inicio", 13), ("fim", 13), ("categoria", 15), ("pavimento", 20), ("inicio_real", 13), ("fim_real", 13), ("avanco", 10)],
    [[e[0], e[1], d(e[2]), d(e[3]), e[4], e[5], d(e[6]), d(e[7]), e[8]] for e in ETAPAS],
    {5: "categoria"},
    datas={3, 4, 7, 8},
)

# ---------- Vínculos ----------
tabela(wb.create_sheet("Vínculos"), [("guid", 26), ("tarefa", 12), ("acao", 14), ("modo", 10)], [], {3: "acao", 4: "modo"})

# ---------- Falas ----------
tabela(
    wb.create_sheet("Falas"),
    [("ordem", 8), ("arquivo", 28), ("assunto", 40), ("cena", 16), ("recorte", 14), ("inicio_s", 10), ("fim_s", 10)],
    [],
    {4: "cena", 5: "recorte"},
)

# ---------- Fotos ----------
tabela(
    wb.create_sheet("Fotos"),
    [("arquivo", 28), ("data", 13), ("local", 22), ("descricao", 44), ("etapa", 10)],
    [],
    {},
    datas={2},
)

# ---------- Vídeo ----------
ficha(wb.create_sheet("Vídeo"), [
    ("Formato", "vertical 9:16", "Vertical para Reels e Stories; horizontal para YouTube e apresentações.", "formato"),
    ("Duração (s)", "", "Vazio = soma das falas (ou 30 s sem falas). Máximo 120 s.", None),
    ("Quadros por segundo", "30", "", "fps"),
    ("Qualidade", "máxima", "Máxima renderiza com mais definição e demora mais.", "qualidade"),
    ("Aparência", "Realista", "", "aparencia"),
    ("Luz", "Dia", "Ciclo = o dia corre do amanhecer à noite durante o voo.", "luz"),
    ("Animação", "Progressivo", "Como cada elemento surge durante a sua etapa.", "animacao"),
    ("Marca no canto", "sim", "Faixa com o monograma, o nome e o slogan.", "simnao"),
])

# ---------- Documento ----------
ficha(wb.create_sheet("Documento"), [
    ("Título", "Relatório de acompanhamento da obra", "", None),
    ("Destinatário", "Família Exemplo", "", None),
    ("Observações", "", "Texto livre no fim do documento.", None),
    ("Ficha da obra", "sim", "", "simnao"),
    ("Etapas", "sim", "Tabela do cronograma, avanço e desvios.", "simnao"),
    ("Linha do tempo em imagens", "sim", "Quadros do vídeo gerado.", "simnao"),
    ("Fotos da obra", "sim", "", "simnao"),
    ("Referência ao vídeo", "sim", "Nome, duração e código de conferência do MP4.", "simnao"),
])

wb.move_sheet("Listas", offset=len(wb.sheetnames))
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = DOURADO if ws.title == "LEIA-ME" else GRAFITE
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print(f"{SAIDA}: {', '.join(wb.sheetnames)}")
